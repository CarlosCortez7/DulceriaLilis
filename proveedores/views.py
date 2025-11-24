from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.http import HttpResponse, JsonResponse 
from django.db.models import Q
from django.core.paginator import Paginator, EmptyPage, PageNotAnInteger
from django.views.decorators.http import require_POST
from django.contrib.auth.decorators import login_required, permission_required 
import openpyxl
import json
from openpyxl.styles import Font, Alignment
from .models import (
    Proveedor, ProductoProveedor, 
    OrdenCompra, DetalleOrdenCompra
)
from productos.models import Producto
from .forms import (
    ProveedorForm, ProductoProveedorForm, 
    OrdenCompraForm, DetalleOrdenCompraForm
)
from productos.forms import ProductoForm

# ==============================================================================
# HELPER: Lógica común para obtener el listado con filtros y memoria
# (Esto evita repetir código en 'modulo_proveedores' y 'editar_proveedor')
# ==============================================================================
def obtener_listado_proveedores(request):
    """
    Retorna un diccionario con page_obj, query, page_size, sort_by, direction
    usando la lógica de Sesiones y GET.
    """
    # 1. Limpieza
    if request.GET.get('limpiar'):
        if 'prov_search' in request.session: del request.session['prov_search']
        if 'prov_page_size' in request.session: del request.session['prov_page_size']
        if 'prov_page' in request.session: del request.session['prov_page']
        return None # Señal para redirigir

    # 2. Búsqueda (Search)
    req_query = request.GET.get('q')
    if req_query is not None:
        query = req_query
        request.session['prov_search'] = query
        request.session['prov_page'] = 1 
    else:
        query = request.session.get('prov_search', '')

    # 3. Tamaño de Página
    req_page_size = request.GET.get('page_size')
    if req_page_size:
        page_size = int(req_page_size)
        request.session['prov_page_size'] = page_size
    else:
        page_size = int(request.session.get('prov_page_size', 10))

    # 4. Número de Página
    req_page = request.GET.get('page')
    if req_page:
        page_number = req_page
        request.session['prov_page'] = page_number
    else:
        page_number = request.session.get('prov_page', 1)

    # 5. Ordenamiento
    sort_by = request.GET.get('sort', 'razon_social')
    direction = request.GET.get('dir', 'asc')
    order_prefix = '-' if direction == 'desc' else ''
    
    allowed_sort_fields = ['razon_social', 'rut_nif', 'estado']
    if sort_by not in allowed_sort_fields: sort_by = 'razon_social'

    # 6. QuerySet
    if query:
        proveedores_list = Proveedor.objects.filter(
            Q(rut_nif__icontains=query) | 
            Q(razon_social__icontains=query) |
            Q(nombre_fantasia__icontains=query)
        ).order_by(f"{order_prefix}{sort_by}")
    else:
        proveedores_list = Proveedor.objects.all().order_by(f"{order_prefix}{sort_by}")

    # 7. Paginación
    paginator = Paginator(proveedores_list, page_size)
    try:
        page_obj = paginator.page(page_number)
    except PageNotAnInteger:
        page_obj = paginator.page(1)
    except EmptyPage:
        page_obj = paginator.page(paginator.num_pages)

    return {
        'proveedores': page_obj, # Legacy support
        'page_obj': page_obj,
        'query': query,
        'page_size': page_size,
        'sort_by': sort_by,
        'direction': direction
    }


# ------------------------------
# VISTA PROVEEDORES (LISTAR + CREAR)
# ------------------------------
@login_required 
@permission_required('proveedores.view_proveedor', raise_exception=True)
def modulo_proveedores(request):
    
    # Usamos el helper
    context_listado = obtener_listado_proveedores(request)
    if context_listado is None: # Si se pidió limpiar
        return redirect('modulo_proveedores')

    if request.method == 'POST':
        if not request.user.has_perm('proveedores.add_proveedor'):
             messages.error(request, "No tienes permisos para crear proveedores.")
             return redirect('modulo_proveedores')

        form = ProveedorForm(request.POST)
        if form.is_valid():
            nuevo_proveedor = form.save() 
            messages.success(request, f"Proveedor '{nuevo_proveedor.razon_social}' guardado.")
            return redirect('editar_proveedor', proveedor_id=nuevo_proveedor.id) 
        else:
            messages.error(request, "Corrige los errores en el formulario.")
            form_para_mostrar = form
    else:
        form_para_mostrar = ProveedorForm()

    context = {
        'titulo_pagina': 'Módulo de Proveedores',
        'form': form_para_mostrar, 
        'editando': False,
    }
    # Fusionamos los datos del listado (page_obj, query, etc) al contexto principal
    context.update(context_listado)
    
    return render(request, 'proveedores/modulo_proveedores.html', context)


# ------------------------------
# VISTA PROVEEDORES (EDICIÓN)
# ------------------------------
@login_required 
@permission_required('proveedores.change_proveedor', raise_exception=True)
def editar_proveedor(request, proveedor_id=None, pk=None):
    
    id_proveedor = proveedor_id or pk
    proveedor_a_editar = get_object_or_404(Proveedor, pk=id_proveedor)

    # --- AQUÍ ESTABA EL ERROR: FALTABA CARGAR EL CONTEXTO DEL LISTADO ---
    context_listado = obtener_listado_proveedores(request)
    if context_listado is None: 
        return redirect('modulo_proveedores')
    # --------------------------------------------------------------------

    if request.method == 'POST':
        form = ProveedorForm(request.POST, instance=proveedor_a_editar)
        if form.is_valid():
            form.save()
            messages.success(request, 'Proveedor actualizado correctamente.')
            return redirect('editar_proveedor', proveedor_id=id_proveedor)
        else:
            messages.error(request, 'Corrige los errores en el formulario.')
            form_para_mostrar = form
    else:
        form_para_mostrar = ProveedorForm(instance=proveedor_a_editar)

    # Lógica Tabs Productos
    productos_asociados = proveedor_a_editar.productos_proveedor.select_related('producto').all()
    ids_productos_asociados = productos_asociados.values_list('producto__id', flat=True)
    productos_disponibles = Producto.objects.exclude(id__in=ids_productos_asociados).order_by('nombre')
    
    pp_form_existente = ProductoProveedorForm(prefix="existente")
    pp_form_nuevo = ProductoProveedorForm(prefix="nuevo")
    initial_data_producto = { 'uom_venta': 'unidad', 'factor_conversion': 1, 'impuesto_iva': 19.00, 'stock_actual': 0, 'stock_minimo': 0, 'perishable': False, 'control_por_lote': False, 'control_por_serie': False, 'estado': 'activo' }
    producto_form = ProductoForm(initial=initial_data_producto) 

    context = {
        'titulo_pagina': f'Editando a: {proveedor_a_editar.razon_social}',
        'form': form_para_mostrar, 
        'editando': True, 
        'proveedor_a_editar': proveedor_a_editar,
        'productos_asociados': productos_asociados,
        'productos_disponibles': productos_disponibles,
        'pp_form_existente': pp_form_existente,
        'pp_form_nuevo': pp_form_nuevo,
        'producto_form': producto_form
    }
    
    # Fusionamos el contexto del listado para que la tabla de abajo funcione
    context.update(context_listado)

    return render(request, 'proveedores/modulo_proveedores.html', context)


# ------------------------------
# VISTAS AUXILIARES (TAB 3 PROVEEDORES)
# ------------------------------
@login_required
@permission_required('proveedores.change_proveedor', raise_exception=True) 
@require_POST
def asociar_producto_proveedor(request, proveedor_id):
    proveedor = get_object_or_404(Proveedor, id=proveedor_id)
    
    if 'submit_nuevo' in request.POST and not request.user.has_perm('productos.add_producto'):
        messages.error(request, "No tienes permisos para crear nuevos productos.")
        return redirect('editar_proveedor', proveedor_id=proveedor_id)

    pp_form = None
    producto_form = None
    
    if 'submit_existente' in request.POST:
        pp_form = ProductoProveedorForm(request.POST, prefix="existente")
        producto_id = request.POST.get('producto') 
        
        if not producto_id:
            messages.error(request, "Error: Debe seleccionar un producto.")
            return redirect('editar_proveedor', proveedor_id=proveedor_id)
            
        producto = get_object_or_404(Producto, id=producto_id)
        
        if pp_form.is_valid():
            if ProductoProveedor.objects.filter(proveedor=proveedor, producto=producto).exists():
                messages.warning(request, f"El producto '{producto.nombre}' ya estaba asociado.")
                return redirect('editar_proveedor', proveedor_id=proveedor_id)
            
            asociacion = pp_form.save(commit=False)
            asociacion.proveedor = proveedor
            asociacion.producto = producto
            asociacion.save()
            messages.success(request, f"Producto '{producto.nombre}' asociado correctamente.")
        else:
            messages.error(request, f"Error en los datos de costo/lead time: {pp_form.errors.as_text()}")

    elif 'submit_nuevo' in request.POST:
        pp_form = ProductoProveedorForm(request.POST, prefix="nuevo")
        producto_form = ProductoForm(request.POST, request.FILES)

        if producto_form.is_valid() and pp_form.is_valid():
            try:
                nuevo_producto = producto_form.save()
                asociacion = pp_form.save(commit=False)
                asociacion.proveedor = proveedor
                asociacion.producto = nuevo_producto
                asociacion.save()
                messages.success(request, f"Producto NUEVO '{nuevo_producto.nombre}' creado y asociado.")
            except Exception as e:
                if 'nuevo_producto' in locals() and nuevo_producto.id:
                    nuevo_producto.delete()
                messages.error(request, f"Error al crear o asociar: {e}")
        else:
            msg = "Error al crear el producto: " + producto_form.errors.as_text() if not producto_form.is_valid() else ""
            msg += " Error en los datos de costo: " + pp_form.errors.as_text() if not pp_form.is_valid() else ""
            messages.error(request, msg)

    return redirect('editar_proveedor', proveedor_id=proveedor_id)


@login_required
@permission_required('proveedores.change_proveedor', raise_exception=True)
@require_POST
def desasociar_producto_proveedor(request, asociacion_id):
    asociacion = get_object_or_404(ProductoProveedor, id=asociacion_id)
    producto_nombre = asociacion.producto.nombre
    try:
        asociacion.delete()
        return JsonResponse({'status': 'success', 'message': f"Producto '{producto_nombre}' desasociado."})
    except Exception as e:
        return JsonResponse({'status': 'error', 'message': f"Error al desasociar: {e}"}, status=400)


# ------------------------------
# ELIMINAR PROVEEDOR
# ------------------------------
@login_required
@permission_required('proveedores.delete_proveedor', raise_exception=True)
@require_POST
def eliminar_proveedor(request, proveedor_id=None, pk=None):
    id_proveedor = proveedor_id or pk
    proveedor = get_object_or_404(Proveedor, pk=id_proveedor)
    proveedor_nombre = proveedor.razon_social
    
    try:
        proveedor.delete()
        return JsonResponse({'status': 'success', 'message': f"Proveedor '{proveedor_nombre}' eliminado exitosamente."})
    except Exception as e:
        msg = f"No se pudo eliminar. Puede tener datos asociados. Error: {e}"
        return JsonResponse({'status': 'error', 'message': msg}, status=400)


# ------------------------------
# EXPORTAR A EXCEL (PROVEEDORES)
# ------------------------------
@login_required
@permission_required('proveedores.view_proveedor', raise_exception=True) 
def exportar_excel_proveedores(request):
    # Reutilizamos la lógica del query de sesión
    query = request.session.get('prov_search', '')
    if query:
        proveedores = Proveedor.objects.filter(
            Q(rut_nif__icontains=query) | 
            Q(razon_social__icontains=query) |
            Q(nombre_fantasia__icontains=query)
        ).order_by('razon_social')
    else:
        proveedores = Proveedor.objects.all().order_by('razon_social')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Proveedores"

    headers = [
        "ID", "Razón Social", "Nombre Fantasía", "RUT/NIF",
        "Contacto Principal", "Email Contacto", "Teléfono Contacto",
        "Email General", "Teléfono General", "Sitio Web",
        "Dirección", "Ciudad", "País", "Condiciones de Pago",
        "Moneda", "Tipo", "Estado", "Observaciones", "Creado", "Actualizado"
    ]
    
    header_font = Font(bold=True)
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.font = header_font
        ws.column_dimensions[cell.column_letter].width = 20
    
    for row_num, p in enumerate(proveedores, 2):
        ws.cell(row=row_num, column=1, value=p.id)
        ws.cell(row=row_num, column=2, value=p.razon_social)
        ws.cell(row=row_num, column=3, value=p.nombre_fantasia)
        ws.cell(row=row_num, column=4, value=p.rut_nif)
        ws.cell(row=row_num, column=5, value=p.contacto_principal_nombre)
        ws.cell(row=row_num, column=6, value=p.contacto_principal_email)
        ws.cell(row=row_num, column=7, value=p.contacto_principal_telefono)
        ws.cell(row=row_num, column=8, value=p.email)
        ws.cell(row=row_num, column=9, value=p.telefono)
        ws.cell(row=row_num, column=10, value=p.sitio_web)
        ws.cell(row=row_num, column=11, value=p.direccion)
        ws.cell(row=row_num, column=12, value=p.ciudad)
        ws.cell(row=row_num, column=13, value=p.pais)
        ws.cell(row=row_num, column=14, value=p.condiciones_pago)
        ws.cell(row=row_num, column=15, value=p.moneda)
        ws.cell(row=row_num, column=16, value=p.get_tipo_display())
        ws.cell(row=row_num, column=17, value=p.get_estado_display())
        ws.cell(row=row_num, column=18, value=p.observaciones)
        ws.cell(row=row_num, column=19, value=p.creado.strftime('%Y-%m-%d %H:%M') if p.creado else '')
        ws.cell(row=row_num, column=20, value=p.actualizado.strftime('%Y-%m-%d %H:%M') if p.actualizado else '')

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = 'attachment; filename=listado_proveedores.xlsx'
    wb.save(response)
    return response

# ------------------------------------------------------------------
# VISTAS DE ÓRDENES DE COMPRA (Maestro-Detalle)
# (Se mantienen igual que tu versión original funcional, solo asegurando imports)
# ------------------------------------------------------------------

@login_required
@permission_required('proveedores.view_ordencompra', raise_exception=True)
def gestion_orden_compra(request, orden_id=None):
    orden_a_editar = None
    detalles_orden = None
    detalle_form = None
    productos_data_json = "{}" 

    if orden_id:
        if not request.user.has_perm('proveedores.change_ordencompra'):
            messages.error(request, "No tienes permisos para editar esta orden.")
            return redirect('gestion_orden_compra')
            
        orden_a_editar = get_object_or_404(OrdenCompra.objects.select_related('proveedor', 'usuario'), pk=orden_id)
        detalles_orden = orden_a_editar.detalles_orden.select_related('producto').all()
        detalle_form = DetalleOrdenCompraForm(proveedor=orden_a_editar.proveedor)
        
        productos_data = {
            pp.id: float(pp.costo) 
            for pp in detalle_form.fields['producto_proveedor'].queryset
        }
        productos_data_json = json.dumps(productos_data)

    if request.method == 'POST':
        if orden_a_editar:
            form = OrdenCompraForm(request.POST, instance=orden_a_editar)
        else:
            if not request.user.has_perm('proveedores.add_ordencompra'):
                messages.error(request, "No tienes permisos para crear órdenes.")
                return redirect('gestion_orden_compra')
            form = OrdenCompraForm(request.POST)

        if form.is_valid():
            orden_guardada = form.save(commit=False)
            if not orden_a_editar:
                orden_guardada.usuario = request.user
            orden_guardada.save()
            
            if orden_a_editar:
                messages.success(request, f"Estado de la Orden #{orden_guardada.id} actualizado.")
                return redirect('editar_orden_compra', orden_id=orden_guardada.id)
            else:
                messages.success(request, f"Orden #{orden_guardada.id} creada. Ahora añada productos.")
                return redirect('editar_orden_compra', orden_id=orden_guardada.id)
        else:
            messages.error(request, "Error al guardar la orden. Revisa el formulario.")
            form_para_mostrar = form
    else:
        if orden_a_editar:
            form_para_mostrar = OrdenCompraForm(instance=orden_a_editar)
        else:
            form_para_mostrar = OrdenCompraForm()

    # Paginador Órdenes
    if 'page_size' in request.GET:
        page_size = request.GET.get('page_size', 10)
        request.session['oc_page_size'] = page_size
    else:
        page_size = request.session.get('oc_page_size', 10)

    query = request.GET.get('q', '')
    sort_by = request.GET.get('sort', 'fecha_emision')
    direction = request.GET.get('dir', 'desc')
    order_prefix = '-' if direction == 'desc' else ''
    
    allowed_sort_fields = ['id', 'proveedor__razon_social', 'estado', 'total', 'fecha_emision']
    if sort_by not in allowed_sort_fields: sort_by = 'fecha_emision'
        
    if query:
        ordenes = OrdenCompra.objects.filter(
            Q(id__icontains=query) | Q(proveedor__razon_social__icontains=query) | Q(estado__icontains=query)
        ).select_related('proveedor', 'usuario').order_by(f"{order_prefix}{sort_by}")
    else:
        ordenes = OrdenCompra.objects.all().select_related('proveedor', 'usuario').order_by(f"{order_prefix}{sort_by}")

    paginator = Paginator(ordenes, page_size)
    page_number = request.GET.get('page')
    try:
        page_obj = paginator.page(page_number)
    except PageNotAnInteger:
        page_obj = paginator.page(1)
    except EmptyPage:
        page_obj = paginator.page(paginator.num_pages)

    context = {
        'titulo_pagina': 'Gestión de Órdenes de Compra',
        'form': form_para_mostrar,
        'page_obj': page_obj,
        'query': query,
        'editando': bool(orden_a_editar),
        'orden_a_editar': orden_a_editar,
        'detalles_orden': detalles_orden,
        'detalle_form': detalle_form,
        'productos_data_json': productos_data_json,
        'page_size': int(page_size), 
        'sort_by': sort_by, 
        'direction': direction 
    }
    return render(request, 'ordencompra/modulodeordenes.html', context)


@login_required
@permission_required('proveedores.change_ordencompra', raise_exception=True) 
@require_POST
def agregar_detalle_orden(request, orden_id):
    orden = get_object_or_404(OrdenCompra, id=orden_id)
    form = DetalleOrdenCompraForm(request.POST, proveedor=orden.proveedor)

    if form.is_valid():
        detalle = form.save(commit=False)
        detalle.orden = orden
        
        producto_proveedor = form.cleaned_data['producto_proveedor']
        detalle.producto = producto_proveedor.producto 
        
        detalle_existente = DetalleOrdenCompra.objects.filter(
            orden=orden, 
            producto=detalle.producto
        ).first()

        if detalle_existente:
            detalle_existente.cantidad += detalle.cantidad
            if detalle_existente.precio_unitario != detalle.precio_unitario:
                detalle_existente.precio_unitario = detalle.precio_unitario
            detalle_existente.save()
            messages.success(request, f"Cantidad de '{detalle.producto.nombre}' actualizada.")
        else:
            detalle.save()
            messages.success(request, f"Producto '{detalle.producto.nombre}' añadido a la orden.")
    else:
        error_msg = f"Error al añadir el producto: {form.errors.as_text()}"
        messages.error(request, error_msg)

    return redirect('editar_orden_compra', orden_id=orden_id)


@login_required
@permission_required('proveedores.change_ordencompra', raise_exception=True) 
@require_POST
def eliminar_detalle_orden(request, detalle_id):
    try:
        detalle = get_object_or_404(DetalleOrdenCompra, id=detalle_id)
        producto_nombre = "Producto"
        if detalle.producto:
            producto_nombre = detalle.producto.nombre
        detalle.delete()
        return JsonResponse({'status': 'success', 'message': f"Producto '{producto_nombre}' quitado de la orden."})
    except Exception as e:
        return JsonResponse({'status': 'error', 'message': f"Error al quitar el producto: {e}"}, status=400)


@login_required
@permission_required('proveedores.delete_ordencompra', raise_exception=True) 
@require_POST
def eliminar_orden_compra(request, orden_id):
    orden = get_object_or_404(OrdenCompra, id=orden_id)
    try:
        orden.delete()
        return JsonResponse({'status': 'success', 'message': f"Orden de Compra #{orden_id} eliminada exitosamente."})
    except Exception as e:
        return JsonResponse({'status': 'error', 'message': f"No se pudo eliminar la orden. Error: {e}"}, status=400)


@login_required
@permission_required('proveedores.view_ordencompra', raise_exception=True)
def exportar_excel_ordenes(request):
    query = request.GET.get('q', '')
    sort_by = request.GET.get('sort', 'fecha_emision')
    direction = request.GET.get('dir', 'desc')
    order_prefix = '-' if direction == 'desc' else ''
    
    if query:
        ordenes = OrdenCompra.objects.filter(
            Q(id__icontains=query) | Q(proveedor__razon_social__icontains=query) | Q(estado__icontains=query)
        ).select_related('proveedor', 'usuario').order_by(f"{order_prefix}{sort_by}")
    else:
        ordenes = OrdenCompra.objects.all().select_related('proveedor', 'usuario').order_by(f"{order_prefix}{sort_by}")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ordenes de Compra"

    headers = ["ID Orden", "Proveedor", "RUT Proveedor", "Fecha Emisión", "Estado", "Total", "Emitida por (Usuario)"]
    
    header_font = Font(bold=True)
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.font = header_font
        ws.column_dimensions[cell.column_letter].width = 25

    for row_num, oc in enumerate(ordenes, 2):
        ws.cell(row=row_num, column=1, value=oc.id)
        ws.cell(row=row_num, column=2, value=oc.proveedor.razon_social if oc.proveedor else 'N/A')
        ws.cell(row=row_num, column=3, value=oc.proveedor.rut_nif if oc.proveedor else 'N/A')
        ws.cell(row=row_num, column=4, value=oc.fecha_emision.strftime('%Y-%m-%d %H:%M') if oc.fecha_emision else '')
        ws.cell(row=row_num, column=5, value=oc.get_estado_display())
        cell_total = ws.cell(row=row_num, column=6, value=oc.total)
        cell_total.number_format = '$#,##0.00' 
        ws.cell(row=row_num, column=7, value=oc.usuario.username if oc.usuario else 'N/A')

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = 'attachment; filename=listado_ordenes_compra.xlsx'
    wb.save(response)
    return response